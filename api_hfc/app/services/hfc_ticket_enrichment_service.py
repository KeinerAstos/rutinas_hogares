from __future__ import annotations

import csv
import re
import threading
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook

_BACKEND = Path(r"C:\xampp\htdocs\rutinas_hogares\api_legacy_8011")
_DATA_DIR = Path(r"C:\ProgramData\CentralNOC\Dashboard_Hogar\data\gestion_incidentes")
_STATE_DIR = Path(r"C:\ProgramData\CentralNOC\Dashboard_Hogar\data\hfc_ticket_enrichment")
_SAFE_GLOB = "ENLACE_SEGURO_NODE_INC_TAS_OT_V46_*.csv"

_BACKLOG = _DATA_DIR / "latest_backlog.xlsx"
_WORKORDERS = _DATA_DIR / "latest_workorders.xlsx"
_TAREAS = _DATA_DIR / "latest_tareas.xlsx"
_ALARMAS = _DATA_DIR / "latest_alarmas.xlsx"

_LOCK = threading.Lock()
_CACHE: Dict[str, Any] = {
    "signature": None,
    "by_node": {},
    "duplicate_nodes": set(),
    "built_at": None,
    "source_safe_csv": None,
}

_CLOSED = (
    "CERRADO","CERRADA","CLOSED",
    "CANCELADO","CANCELADA","CANCELLED",
    "RESUELTO","RESUELTA","RESOLVED",
    "COMPLETADO","COMPLETADA","COMPLETED",
    "FINALIZADO","FINALIZADA",
)

def _norm(v: Any) -> str:
    if v is None:
        return ""
    s = str(v).strip().upper().replace("\u00A0", " ")
    return re.sub(r"\s+", " ", s)

def _parse_dt(v: Any) -> Optional[datetime]:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v
    s = str(v).strip().replace("T", " ")
    s = re.sub(r"Z$", "", s)
    s = re.sub(r"([+-]\d\d:\d\d)$", "", s)
    for fmt in (
        "%Y-%m-%d %H:%M:%S","%Y-%m-%d %H:%M",
        "%d/%m/%Y %H:%M:%S","%d/%m/%Y %H:%M",
        "%d-%m-%Y %H:%M:%S","%d-%m-%Y %H:%M",
        "%Y/%m/%d %H:%M:%S","%Y-%m-%d","%d/%m/%Y","%d-%m-%Y",
    ):
        try:
            return datetime.strptime(s, fmt)
        except Exception:
            pass
    try:
        return datetime.fromisoformat(s)
    except Exception:
        return None

def _active(v: Any) -> Optional[bool]:
    s = _norm(v)
    if not s:
        return None
    return not any(x in s for x in _CLOSED)

def _age_days(dt: Optional[datetime]) -> Optional[float]:
    if not dt:
        return None
    return round((datetime.now() - dt).total_seconds() / 86400.0, 2)

def _read_sheet(path: Path, sheet: str = "Exportar Hoja de Trabajo") -> Tuple[List[Dict[str, Any]], List[str]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise RuntimeError(f"Hoja {sheet!r} no existe en {path}")
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        headers = [str(x).strip() if x is not None else "" for x in (header or [])]
        rows: List[Dict[str, Any]] = []
        for vals in it:
            row: Dict[str, Any] = {}
            nonempty = False
            for i, h in enumerate(headers):
                if not h:
                    continue
                v = vals[i] if i < len(vals) else None
                row[h] = v
                if v not in (None, ""):
                    nonempty = True
            if nonempty:
                rows.append(row)
        return rows, headers
    finally:
        wb.close()

def _pick(row: Dict[str, Any], *names: str) -> Any:
    cmap = {_norm(k): k for k in row.keys()}
    for n in names:
        k = cmap.get(_norm(n))
        if k is not None:
            return row.get(k)
    return None

def _latest_safe_csv() -> Path:
    files = sorted(_STATE_DIR.glob(_SAFE_GLOB), key=lambda p: p.stat().st_mtime, reverse=True)
    if not files:
        raise RuntimeError("No existe ENLACE_SEGURO_NODE_INC_TAS_OT_V46_*.csv")
    return files[0]

def _load_safe_nodes(path: Path) -> Tuple[set[str], set[str]]:
    safe: set[str] = set()
    duplicates: set[str] = set()
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        reader = csv.DictReader(fh, delimiter=";")
        for row in reader:
            node = _norm(row.get("NODO"))
            if node and _norm(row.get("MATCH_CALIDAD")) == "EXACTO_UNICO":
                safe.add(node)

    amb_files = sorted(
        _STATE_DIR.glob("AMBIGUEDADES_NODE_INC_V46_*.csv"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if amb_files:
        with amb_files[0].open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
            reader = csv.DictReader(fh, delimiter=";")
            for row in reader:
                if _norm(row.get("TIPO")) == "NODE_CANONICO_DUPLICADO":
                    n = _norm(row.get("NODE"))
                    if n:
                        duplicates.add(n)
    return safe, duplicates

def _signature() -> Tuple[Any, ...]:
    safe = _latest_safe_csv()
    paths = [_BACKLOG, _WORKORDERS, _TAREAS, _ALARMAS, safe]
    sig: List[Any] = []
    for p in paths:
        st = p.stat()
        sig.extend([str(p), st.st_mtime_ns, st.st_size])
    return tuple(sig)

def _build_cache() -> Dict[str, Any]:
    safe_csv = _latest_safe_csv()
    safe_nodes, duplicates = _load_safe_nodes(safe_csv)

    for p in (_BACKLOG, _WORKORDERS, _TAREAS, _ALARMAS):
        if not p.is_file():
            raise RuntimeError(f"Fuente no encontrada: {p}")

    backlog, _ = _read_sheet(_BACKLOG)
    ots, _ = _read_sheet(_WORKORDERS)
    tasks, _ = _read_sheet(_TAREAS)
    alarms, _ = _read_sheet(_ALARMAS)

    inc_master: Dict[str, Dict[str, Any]] = {}
    for r in backlog:
        inc = _norm(_pick(r, "INCIDENTE"))
        if inc:
            inc_master[inc] = r

    node_to_inc: Dict[str, set[str]] = {}
    for r in alarms:
        node = _norm(_pick(r, "NODE"))
        inc = _norm(_pick(r, "R_TICKETNRO", "INCIDENTE"))
        if node and inc:
            node_to_inc.setdefault(node, set()).add(inc)

    inc_to_tas: Dict[str, List[Dict[str, Any]]] = {}
    for r in tasks:
        inc = _norm(_pick(r, "INC", "INCIDENTE", "ROOTREQUESTNAME"))
        if inc:
            inc_to_tas.setdefault(inc, []).append(r)

    inc_to_ot: Dict[str, List[Dict[str, Any]]] = {}
    for r in ots:
        inc = _norm(_pick(r, "INCIDENTE", "ROOT_INCIDENT"))
        if inc:
            inc_to_ot.setdefault(inc, []).append(r)

    by_node: Dict[str, Dict[str, Any]] = {}

    for node in sorted(safe_nodes):
        if node in duplicates:
            continue
        incs = sorted(node_to_inc.get(node, set()))
        if not incs:
            continue

        candidates = []
        for inc in incs:
            meta = inc_master.get(inc, {})
            st = _pick(meta, "ESTADO", "STATUS_INC")
            dt = _parse_dt(_pick(meta, "FECHA_CREACION_INC"))
            candidates.append((inc, _active(st), dt, meta))

        active_inc = [x for x in candidates if x[1] is True]
        pool = active_inc or candidates
        pool.sort(key=lambda x: x[2] or datetime.min, reverse=True)
        inc, inc_is_active, inc_dt, meta = pool[0]

        tas_all = inc_to_tas.get(inc, [])
        ot_all = inc_to_ot.get(inc, [])
        tas_active = [r for r in tas_all if _active(_pick(r, "STATUS_TAS", "STATUS")) is True]
        ot_active = [r for r in ot_all if _active(_pick(r, "ESTADO_OT", "STATUS")) is True]

        def newest(rows: Iterable[Dict[str, Any]], field: str) -> List[Dict[str, Any]]:
            return sorted(rows, key=lambda r: _parse_dt(_pick(r, field)) or datetime.min, reverse=True)

        tas_pool = newest(tas_active or tas_all, "FECHA_CREACION_TAS")
        ot_pool = newest(ot_active or ot_all, "FECHA_CREACION_OT")
        tas = tas_pool[0] if tas_pool else None
        ot = ot_pool[0] if ot_pool else None

        tas_status = _pick(tas, "STATUS_TAS", "STATUS") if tas else None
        tas_dt = _parse_dt(_pick(tas, "FECHA_CREACION_TAS")) if tas else None
        ot_status = _pick(ot, "ESTADO_OT", "STATUS") if ot else None
        ot_dt = _parse_dt(_pick(ot, "FECHA_CREACION_OT")) if ot else None

        dias_tk = None
        if inc_is_active is True and inc_dt:
            dias_tk = _age_days(inc_dt)
        elif tas and _active(tas_status) is True and tas_dt:
            dias_tk = _age_days(tas_dt)
        elif ot and _active(ot_status) is True and ot_dt:
            dias_tk = _age_days(ot_dt)

        by_node[node] = {
            "incident": inc,
            "incident_status": _pick(meta, "ESTADO", "STATUS_INC"),
            "incident_opened_at": _pick(meta, "FECHA_CREACION_INC"),
            "tas": _pick(tas, "TAS") if tas else None,
            "tas_status": tas_status,
            "tas_opened_at": _pick(tas, "FECHA_CREACION_TAS") if tas else None,
            "ot": _pick(ot, "WORKORDER") if ot else None,
            "ot_status": ot_status,
            "ot_opened_at": _pick(ot, "FECHA_CREACION_OT") if ot else None,
            "dias_tk_abierto": dias_tk,
            "incidentes_relacionados": len(incs),
            "ticket_ambiguous": len(incs) > 1,
            "ticket_match_quality": "EXACTO_UNICO",
            "ticket_source": "GESTION_INCIDENTES_LATEST",
        }

    return {
        "signature": _signature(),
        "by_node": by_node,
        "duplicate_nodes": duplicates,
        "built_at": datetime.now().isoformat(),
        "source_safe_csv": str(safe_csv),
    }

def _ensure_cache() -> Dict[str, Any]:
    sig = _signature()
    with _LOCK:
        if _CACHE.get("signature") != sig:
            built = _build_cache()
            _CACHE.clear()
            _CACHE.update(built)
        return _CACHE

def enrich_node_record(row: Dict[str, Any]) -> Dict[str, Any]:
    if not isinstance(row, dict):
        return row
    node = _norm(row.get("nodo") or row.get("node"))
    out = dict(row)

    defaults = {
        "incident": None,
        "incident_status": None,
        "incident_opened_at": None,
        "tas": None,
        "tas_status": None,
        "tas_opened_at": None,
        "ot": None,
        "ot_status": None,
        "ot_opened_at": None,
        "dias_tk_abierto": None,
        "incidentes_relacionados": 0,
        "ticket_ambiguous": False,
        "ticket_match_quality": None,
        "ticket_source": None,
    }
    for k, v in defaults.items():
        out.setdefault(k, v)

    if not node:
        return out

    cache = _ensure_cache()
    if node in cache.get("duplicate_nodes", set()):
        out["ticket_match_quality"] = "AMBIGUO_NODO_CANONICO"
        return out

    extra = cache.get("by_node", {}).get(node)
    if extra:
        out.update(extra)
    return out

def enrich_nodes_payload(payload: Any) -> Any:
    if not isinstance(payload, dict):
        return payload

    out = dict(payload)
    for key in ("items", "data", "rows"):
        rows = out.get(key)
        if isinstance(rows, list):
            out[key] = [enrich_node_record(x) if isinstance(x, dict) else x for x in rows]
            out["ticket_enrichment"] = {
                "ok": True,
                "source": "GESTION_INCIDENTES_LATEST",
                "safe_only": True,
            }
            return out
    return out

def enrichment_health() -> Dict[str, Any]:
    cache = _ensure_cache()
    return {
        "ok": True,
        "nodes_enriched": len(cache.get("by_node", {})),
        "duplicate_nodes_blocked": len(cache.get("duplicate_nodes", set())),
        "built_at": cache.get("built_at"),
        "safe_csv": cache.get("source_safe_csv"),
    }